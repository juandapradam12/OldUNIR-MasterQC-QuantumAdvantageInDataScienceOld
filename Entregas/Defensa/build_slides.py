#!/usr/bin/env python3
"""TFM defense deck v3 — complete tribunal narrative, ≤20 slides, image-led."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
FIGS = ROOT / "POCs" / "data" / "output" / "figures"
IMG = ROOT / "Documento" / "Imagenes"
OUT = Path(__file__).resolve().parent / "Defensa_TFM_QuantumAdvantage.pptx"

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.55)

INK = RGBColor(0x14, 0x18, 0x1C)
PAPER = RGBColor(0xFB, 0xFA, 0xF7)
PANEL = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x6B, 0x73, 0x7A)
LINE = RGBColor(0xE6, 0xE1, 0xD8)
ACCENT = RGBColor(0x0B, 0x6E, 0x6A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x0E, 0x12, 0x16)


def _font(run, size, bold=False, color=INK, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _back(slide, shape):
    tree = slide.shapes._spTree
    el = shape._element
    tree.remove(el)
    tree.insert(2, el)


def bg(slide, color=PAPER):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    _fill(sh, color)
    _back(slide, sh)


def rect(slide, left, top, w, h, color, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    return sh


def txt(slide, left, top, w, h, text, size=18, bold=False, color=INK,
        align=PP_ALIGN.LEFT, name="Calibri"):
    box = slide.shapes.add_textbox(left, top, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_after = Pt(3)
        if p.runs:
            _font(p.runs[0], size=size, bold=bold, color=color, name=name)
    return box


def footer(slide, n, total=20):
    txt(slide, M, Inches(7.18), Inches(9), Inches(0.28),
        "UNIR · Máster en Computación Cuántica · Defensa TFM", size=10, color=MUTED)
    txt(slide, Inches(10.5), Inches(7.18), Inches(2.3), Inches(0.28),
        f"{n} / {total}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def header(slide, kicker, title):
    txt(slide, M, Inches(0.28), Inches(12), Inches(0.28),
        kicker.upper(), size=11, bold=True, color=ACCENT)
    txt(slide, M, Inches(0.52), Inches(12.2), Inches(0.55),
        title, size=24, bold=True, color=INK)
    rect(slide, M, Inches(1.12), Inches(1.35), Inches(0.05), ACCENT)


def fit_pic(slide, path: Path, left, top, box_w, box_h):
    rect(slide, left, top, box_w, box_h, PANEL, line=LINE)
    if not path.exists():
        txt(slide, left, top + box_h / 2, box_w, Inches(0.4),
            f"Missing: {path.name}", size=12, color=MUTED, align=PP_ALIGN.CENTER)
        return
    pic = slide.shapes.add_picture(str(path), left, top)
    scale = min(int(box_w) / pic.width, int(box_h) / pic.height)
    nw, nh = int(pic.width * scale), int(pic.height * scale)
    pic.width, pic.height = nw, nh
    pic.left = int(left + (int(box_w) - nw) / 2)
    pic.top = int(top + (int(box_h) - nh) / 2)


def chip(slide, left, top, w, h, value, label):
    rect(slide, left, top, w, h, PANEL, line=LINE)
    rect(slide, left, top, Inches(0.08), h, ACCENT)
    txt(slide, left + Inches(0.22), top + Inches(0.22), w - Inches(0.35), Inches(0.5),
        value, size=20, bold=True, color=ACCENT)
    txt(slide, left + Inches(0.22), top + Inches(0.75), w - Inches(0.35), Inches(0.7),
        label, size=12, color=MUTED)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]
    total = 20
    n = 0

    def new():
        nonlocal n
        n += 1
        return prs.slides.add_slide(blank)

    # 1 Cover
    s = new()
    bg(s, DARK)
    rect(s, 0, 0, Inches(0.22), H, ACCENT)
    if (IMG / "logo_unir.png").exists():
        s.shapes.add_picture(str(IMG / "logo_unir.png"), Inches(0.7), Inches(0.5), height=Inches(0.5))
    txt(s, Inches(0.7), Inches(1.7), Inches(11.5), Inches(0.35),
        "TRABAJO DE FIN DE MÁSTER  ·  DEFENSA", size=12, bold=True, color=ACCENT)
    txt(s, Inches(0.7), Inches(2.2), Inches(11.8), Inches(2.0),
        "Ventaja cuántica en ciencia de datos:\ncondiciones empíricas, métodos híbridos\ny criterios de activación selectiva",
        size=28, bold=True, color=WHITE)
    txt(s, Inches(0.7), Inches(4.8), Inches(11.5), Inches(0.9),
        "Juan David Prada Malagón\nDir. Ricardo Rafael Domínguez Reyes  ·  UNIR  ·  Máster en Computación Cuántica",
        size=14, color=RGBColor(0xA8, 0xB0, 0xB8))
    rect(s, Inches(0.7), Inches(6.3), Inches(2.8), Inches(0.05), ACCENT)
    txt(s, Inches(0.7), Inches(6.5), Inches(11.5), Inches(0.4),
        "9 casos  ·  kernels ZZ  ·  QAOA/VQE  ·  Q(D)  ·  QCR", size=13, color=RGBColor(0x8A, 0x93, 0x9B))

    # 2 Agenda
    s = new()
    bg(s)
    header(s, "Hoja de ruta", "Agenda de la defensa")
    agenda = [
        ("01", "Problema central y objetivos"),
        ("02", "Estado del arte"),
        ("03", "Metodología y algoritmos"),
        ("04", "Preguntas originales (tras la evidencia)"),
        ("05", "Resultados empíricos"),
        ("06", "Limitaciones, aportes y futuro"),
    ]
    for i, (num, label) in enumerate(agenda):
        col, row = i % 3, i // 3
        left = M + col * Inches(4.15)
        top = Inches(1.7) + row * Inches(2.3)
        rect(s, left, top, Inches(3.95), Inches(2.0), PANEL, line=LINE)
        txt(s, left + Inches(0.3), top + Inches(0.45), Inches(3.3), Inches(0.5), num, size=28, bold=True, color=ACCENT)
        txt(s, left + Inches(0.3), top + Inches(1.1), Inches(3.3), Inches(0.6), label, size=16, color=INK)
    footer(s, n, total)

    # 3 Motivation
    s = new()
    bg(s)
    header(s, "Contexto", "Brecha entre teoría cuántica y evidencia aplicable")
    cols = [
        ("Promesa", "Kernels y QAOA/VQE con\naceleración teórica"),
        ("Era NISQ", "Ruido, qubits limitados,\nsimulación costosa"),
        ("Hueco", "Pocos protocolos justos\nen problemas de industria"),
    ]
    for i, (t, b) in enumerate(cols):
        left = M + i * Inches(4.15)
        rect(s, left, Inches(1.45), Inches(3.95), Inches(2.1), PANEL, line=LINE)
        rect(s, left, Inches(1.45), Inches(3.95), Inches(0.08), ACCENT)
        txt(s, left + Inches(0.25), Inches(1.75), Inches(3.45), Inches(0.4), t, size=16, bold=True, color=ACCENT)
        txt(s, left + Inches(0.25), Inches(2.3), Inches(3.45), Inches(1.0), b, size=14, color=INK)
    fit_pic(s, FIGS / "07_complejidad_cuantico_vs_clasico.png", M, Inches(3.8), Inches(12.2), Inches(3.1))
    footer(s, n, total)

    # 4 Central research problem (Cap. 1 — planteamiento)
    s = new()
    bg(s)
    header(s, "Planteamiento del problema", "Problema central del TFM")
    rect(s, M, Inches(1.45), Inches(12.2), Inches(2.5), PANEL, line=LINE)
    rect(s, M, Inches(1.45), Inches(0.1), Inches(2.5), ACCENT)
    txt(s, Inches(1.0), Inches(1.7), Inches(11.2), Inches(2.0),
        "¿En qué condiciones, y para qué tipos de problemas de ciencia de datos,\n"
        "los algoritmos cuánticos disponibles en la era NISQ ofrecen una\n"
        "ventaja real y medible frente a los algoritmos clásicos?",
        size=18, bold=True, color=INK)
    txt(s, Inches(0.7), Inches(4.2), Inches(12.2), Inches(0.35),
        "CÓMO SE ATACA  ·  estudio comparativo empírico", size=12, bold=True, color=ACCENT)
    attacks = [
        ("Clasificación", "Kernels QSVM\nvs SVM / boosting"),
        ("Regresión", "QK-SVR\nvs Ridge / árboles"),
        ("Optimización", "QAOA / VQE\nvs greedy / ILP"),
        ("Híbrido", "QAOA–COBYLA\n(MDS, QUBO)"),
    ]
    for i, (t, d) in enumerate(attacks):
        left = M + i * Inches(3.15)
        rect(s, left, Inches(4.65), Inches(3.0), Inches(1.9), PANEL, line=LINE)
        txt(s, left + Inches(0.2), Inches(4.85), Inches(2.6), Inches(0.4), t, size=14, bold=True, color=ACCENT)
        txt(s, left + Inches(0.2), Inches(5.35), Inches(2.6), Inches(0.9), d, size=13, color=INK)
    footer(s, n, total)

    # 5 Objectives
    s = new()
    bg(s)
    header(s, "Objetivos", "Qué se propuso y se entrega")
    rows = [
        ("01", "Comparar", "Cuántico vs clásico en 9 casos bajo protocolos justos"),
        ("02", "Caracterizar", "Patrón híbrido QAOA/VQE + criterios de activación"),
        ("03", "Diagnosticar", "Score de cuanticidad Q(D) previo al entrenamiento"),
        ("04", "Plantear", "Framework QCR (espacio de representación) — formalmente"),
    ]
    for i, (num, title, desc) in enumerate(rows):
        top = Inches(1.45) + i * Inches(1.25)
        rect(s, M, top, Inches(12.2), Inches(1.1), PANEL, line=LINE)
        txt(s, Inches(0.85), top + Inches(0.28), Inches(1.0), Inches(0.5), num, size=22, bold=True, color=ACCENT)
        txt(s, Inches(2.1), top + Inches(0.22), Inches(3.5), Inches(0.4), title, size=18, bold=True, color=INK)
        txt(s, Inches(5.5), top + Inches(0.28), Inches(6.8), Inches(0.55), desc, size=15, color=MUTED)
    footer(s, n, total)

    # 6 State of the art / positioning
    s = new()
    bg(s)
    header(s, "Estado del arte", "Posicionamiento del TFM")
    pillars = [
        ("Kernels cuánticos", "Schuld, Huang, Liu et al.\nVentaja posible si hay\nestructura geométrica alineada"),
        ("Optimización variacional", "Farhi QAOA · Peruzzo VQE\nParadigma híbrido NISQ\nQUBO → Hamiltoniano"),
        ("Este trabajo", "Protocolos justos + datos SKG\nQ(D) como diagnóstico\nQCR formalmente planteado"),
    ]
    for i, (t, b) in enumerate(pillars):
        left = M + i * Inches(4.15)
        rect(s, left, Inches(1.55), Inches(3.95), Inches(4.8), PANEL, line=LINE)
        rect(s, left, Inches(1.55), Inches(3.95), Inches(0.7), ACCENT if i < 2 else DARK)
        txt(s, left + Inches(0.25), Inches(1.7), Inches(3.45), Inches(0.45), t, size=15, bold=True, color=WHITE)
        txt(s, left + Inches(0.3), Inches(2.7), Inches(3.4), Inches(3.2), b, size=16, color=INK)
    footer(s, n, total)

    # 7 Methodology
    s = new()
    bg(s)
    header(s, "Metodología", "9 casos · 3 capas de datos · comparación emparejada")
    fit_pic(s, FIGS / "07_dashboard_ventaja_cuantica.png", M, Inches(1.35), Inches(12.2), Inches(5.5))
    footer(s, n, total)

    # 8 Design + heatmap
    s = new()
    bg(s)
    header(s, "Diseño experimental", "Qué se midió y dónde aparece ventaja")
    fit_pic(s, FIGS / "07_heatmap_ventaja_por_dominio.png", M, Inches(1.35), Inches(7.5), Inches(5.5))
    side = [
        ("Clasificación", "QSVM vs boosting / RBF"),
        ("Regresión", "QK-SVR vs Ridge / árboles"),
        ("Optimización", "QAOA/VQE vs greedy / ILP"),
        ("Anomalías / lags", "OCSVM-ZZ · QK-SVR"),
        ("Capas de datos", "Nativo · SKG · público"),
    ]
    for i, (t, d) in enumerate(side):
        top = Inches(1.4) + i * Inches(1.05)
        rect(s, Inches(8.3), top, Inches(4.45), Inches(0.95), PANEL, line=LINE)
        txt(s, Inches(8.55), top + Inches(0.15), Inches(4.0), Inches(0.3), t, size=14, bold=True, color=ACCENT)
        txt(s, Inches(8.55), top + Inches(0.48), Inches(4.0), Inches(0.35), d, size=12, color=INK)
    footer(s, n, total)

    # 9 Algorithms kernels
    s = new()
    bg(s)
    header(s, "Algoritmos", "Kernels cuánticos")
    fit_pic(s, FIGS / "06_zzfeaturemap_circuit.png", M, Inches(1.35), Inches(7.4), Inches(5.5))
    steps = [
        ("Codificación", "ZZFeatureMap · 2q · reps=2"),
        ("Kernel", "κ = |⟨φ(xᵢ)|φ(xⱼ)⟩|²"),
        ("Modelos", "QSVM · QK-SVR · OCSVM-ZZ"),
        ("Control justo", "RBF · mismo C · mismo PCA"),
        ("Principio", "Ventaja si mapa ↔ generador"),
    ]
    for i, (t, d) in enumerate(steps):
        top = Inches(1.4) + i * Inches(1.05)
        rect(s, Inches(8.2), top, Inches(4.55), Inches(0.95), PANEL, line=LINE)
        txt(s, Inches(8.45), top + Inches(0.15), Inches(4.1), Inches(0.3), t, size=14, bold=True, color=ACCENT)
        txt(s, Inches(8.45), top + Inches(0.48), Inches(4.1), Inches(0.35), d, size=12, color=INK)
    footer(s, n, total)

    # 10 Algorithms hybrid
    s = new()
    bg(s)
    header(s, "Algoritmos", "Híbridos QAOA / VQE + COBYLA")
    fit_pic(s, FIGS / "10_qaoa_convergencia.png", M, Inches(1.35), Inches(6.0), Inches(5.5))
    fit_pic(s, FIGS / "10_mds_grafo_solucion.png", Inches(6.7), Inches(1.35), Inches(6.0), Inches(5.5))
    footer(s, n, total)

    # 11 Original research questions — wording from investigacion_original.tex
    s = new()
    bg(s)
    header(s, "Preguntas originales", "Surgen tras la evidencia de ventaja condicionada")
    txt(s, M, Inches(1.28), Inches(12.2), Inches(0.5),
        "Hallazgo previo: la ventaja en kernels aparece con alineación mapa↔generador y desaparece sin ella.\n"
        "Eso abre dos preguntas nuevas (Cap. investigación original):",
        size=12, color=MUTED)

    rect(s, M, Inches(1.95), Inches(12.2), Inches(2.25), PANEL, line=LINE)
    rect(s, M, Inches(1.95), Inches(0.08), Inches(2.25), ACCENT)
    txt(s, Inches(0.85), Inches(2.1), Inches(11.5), Inches(0.3),
        "P1  ·  TESTS DE CUANTICIDAD", size=12, bold=True, color=ACCENT)
    txt(s, Inches(0.85), Inches(2.5), Inches(11.5), Inches(1.0),
        "¿Existe algún indicador calculable a priori sobre un dataset que permita\n"
        "predecir si un algoritmo cuántico mostrará ventaja sobre los clásicos?",
        size=15, bold=True, color=INK)
    txt(s, Inches(0.85), Inches(3.55), Inches(11.5), Inches(0.4),
        "Respuesta propuesta en este TFM: score compuesto Q(D)  (alineación · entropía espectral · d_F)",
        size=12, color=MUTED)

    rect(s, M, Inches(4.25), Inches(12.2), Inches(2.5), PANEL, line=LINE)
    rect(s, M, Inches(4.25), Inches(0.08), Inches(2.5), DARK)
    txt(s, Inches(0.85), Inches(4.4), Inches(11.5), Inches(0.3),
        "P2  ·  MAPEOS CUÁNTICO-COMPATIBLES", size=12, bold=True, color=DARK)
    txt(s, Inches(0.85), Inches(4.8), Inches(11.5), Inches(1.0),
        "Si un dataset no tiene estructura cuántica nativa, ¿puede transformarse a un\n"
        "espacio de representación donde la tenga, sin alterar la información discriminativa?",
        size=15, bold=True, color=INK)
    txt(s, Inches(0.85), Inches(5.95), Inches(11.5), Inches(0.5),
        "Respuesta propuesta: framework QCR + conjetura del espacio Z (ψ polinómica) — planteada formalmente, abierta",
        size=12, color=MUTED)
    footer(s, n, total)

    # 12 Results alignment
    s = new()
    bg(s)
    header(s, "Resultados", "Principio de alineación")
    fit_pic(s, FIGS / "06_quantum_kernel_boundaries.png", M, Inches(1.35), Inches(6.0), Inches(5.5))
    fit_pic(s, FIGS / "06_learning_curves_quantum_advantage.png", Inches(6.7), Inches(1.35), Inches(6.0), Inches(5.5))
    footer(s, n, total)

    # 13 Classification + optimization metrics
    s = new()
    bg(s)
    header(s, "Resultados", "Clasificación y optimización — cifras ancla")
    fit_pic(s, FIGS / "01_escalabilidad_ventaja_clasificacion.png", M, Inches(1.35), Inches(7.6), Inches(5.5))
    chip(s, Inches(8.4), Inches(1.4), Inches(4.35), Inches(1.5), "+38.3 pp", "QSVM en datos alineados")
    chip(s, Inches(8.4), Inches(3.1), Inches(4.35), Inches(1.5), "−30.3%", "QAOA vs Greedy (asignación)")
    chip(s, Inches(8.4), Inches(4.8), Inches(4.35), Inches(1.5), "Empate TSP", "QAOA = NN+2-opt (modelado)")
    footer(s, n, total)

    # 14 Anomalies
    s = new()
    bg(s)
    header(s, "Resultados", "Anomalías one-class y forecasting")
    fit_pic(s, IMG / "zz_phase_concept_diagram.png", M, Inches(1.35), Inches(6.0), Inches(4.0))
    fit_pic(s, IMG / "zz_phase_scores_if_vs_zz.png", Inches(6.7), Inches(1.35), Inches(6.0), Inches(4.0))
    # bottom strip findings
    rect(s, M, Inches(5.55), Inches(6.0), Inches(1.25), PANEL, line=LINE)
    txt(s, Inches(0.75), Inches(5.7), Inches(5.6), Inches(0.9),
        "zz_phase: OCSVM-ZZ AUC 0.864 vs IF 0.553\nmoons/blobs/spikes: gana el clásico",
        size=13, color=INK)
    rect(s, Inches(6.7), Inches(5.55), Inches(6.0), Inches(1.25), PANEL, line=LINE)
    txt(s, Inches(6.95), Inches(5.7), Inches(5.6), Inches(0.9),
        "zz_lag: QK-SVR MAE 0.056 vs Ridge 0.414\nsine / skg_speed: gana el clásico",
        size=13, color=INK)
    footer(s, n, total)

    # 15 Q(D)
    s = new()
    bg(s)
    header(s, "Respuesta a P1", "Score Q(D): indicador a priori de ventaja")
    fit_pic(s, FIGS / "08_correlacion_indicadores_ventaja.png", M, Inches(1.35), Inches(6.0), Inches(5.5))
    fit_pic(s, FIGS / "08_barrido_sintetico_alpha.png", Inches(6.7), Inches(1.35), Inches(6.0), Inches(5.5))
    footer(s, n, total)

    # 16 QCR
    s = new()
    bg(s)
    header(s, "Respuesta a P2", "QCR: ψ : X → Z  ·  conjetura abierta")
    fit_pic(s, FIGS / "09_qcr_framework_diagnostico.png", M, Inches(1.35), Inches(6.0), Inches(4.6))
    fit_pic(s, FIGS / "09_mapeos_comparativa.png", Inches(6.7), Inches(1.35), Inches(6.0), Inches(4.6))
    rect(s, M, Inches(6.1), Inches(12.2), Inches(0.75), PANEL, line=LINE)
    txt(s, Inches(0.75), Inches(6.22), Inches(11.8), Inches(0.5),
        "Alcance del TFM: plantear formalmente la conjetura y aportar evidencia preliminar — no demostración.",
        size=12, color=MUTED)
    footer(s, n, total)

    # 17 Synthesis
    s = new()
    bg(s)
    header(s, "Síntesis", "Mensaje para el tribunal")
    fit_pic(s, FIGS / "07_radar_ventaja_cuantica.png", M, Inches(1.35), Inches(7.4), Inches(5.5))
    points = [
        "Ventaja real pero condicionada",
        "Kernels: solo con alineación",
        "Híbridos: valor en QUBO moderado",
        "Q(D): diagnóstico útil (ρ_S=0.833)",
        "QCR: planteado, no demostrado",
        "Adopción: selectiva e híbrida",
    ]
    for i, p in enumerate(points):
        top = Inches(1.45) + i * Inches(0.85)
        rect(s, Inches(8.2), top, Inches(4.55), Inches(0.75), PANEL, line=LINE)
        txt(s, Inches(8.45), top + Inches(0.2), Inches(4.1), Inches(0.4), p, size=13, color=INK)
    footer(s, n, total)

    # 18 Limitations
    s = new()
    bg(s)
    header(s, "Limitaciones", "Lo que este TFM no cierra")
    lims = [
        ("Simulación", "Statevector exacto — sin ruido NISQ de hardware real"),
        ("Dimensionalidad", "ZZ 2 qubits → PCA agresivo (~68% varianza fuera)"),
        ("QCR", "Formalmente planteado; validación fuerte = post-TFM"),
        ("n en datos reales", "Correlación Q(D) orientativa; ancla = barrido sintético"),
        ("TSP", "Empate óptimo ≠ speedup; valor de modelado QUBO"),
    ]
    for i, (t, d) in enumerate(lims):
        top = Inches(1.45) + i * Inches(1.0)
        rect(s, M, top, Inches(12.2), Inches(0.9), PANEL, line=LINE)
        txt(s, Inches(0.85), top + Inches(0.25), Inches(2.8), Inches(0.4), t, size=15, bold=True, color=ACCENT)
        txt(s, Inches(3.9), top + Inches(0.25), Inches(8.5), Inches(0.45), d, size=14, color=INK)
    footer(s, n, total)

    # 19 Contributions + future
    s = new()
    bg(s)
    header(s, "Aportes y futuro", "Qué queda cerrado / qué sigue")
    rect(s, M, Inches(1.45), Inches(6.0), Inches(5.3), PANEL, line=LINE)
    rect(s, M, Inches(1.45), Inches(6.0), Inches(0.65), ACCENT)
    txt(s, Inches(0.85), Inches(1.58), Inches(5.4), Inches(0.4), "APORTES DE ESTE TFM", size=14, bold=True, color=WHITE)
    txt(s, Inches(0.85), Inches(2.4), Inches(5.4), Inches(4.0),
        "• Mapa empírico en 9 casos\n"
        "• Protocolos justos reproducibles\n"
        "• Score Q(D) con evidencia sintética\n"
        "• QCR formalmente planteado\n"
        "• Arquitectura híbrida selectiva\n"
        "• Delimitación honesta de no-ventaja",
        size=16, color=INK)

    rect(s, Inches(6.8), Inches(1.45), Inches(6.0), Inches(5.3), PANEL, line=LINE)
    rect(s, Inches(6.8), Inches(1.45), Inches(6.0), Inches(0.65), DARK)
    txt(s, Inches(7.1), Inches(1.58), Inches(5.4), Inches(0.4), "LÍNEAS POST-TFM", size=14, bold=True, color=WHITE)
    txt(s, Inches(7.1), Inches(2.4), Inches(5.4), Inches(4.0),
        "• Hardware ruidoso / shots reales\n"
        "• Más qubits / menos PCA\n"
        "• Validación fuerte de QCR\n"
        "• Q-means / autoencoders\n"
        "• Benchmarks con separación formal\n"
        "• Integración operativa selectiva",
        size=16, color=INK)
    footer(s, n, total)

    # 20 Close
    s = new()
    bg(s, DARK)
    rect(s, 0, 0, Inches(0.22), H, ACCENT)
    txt(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(0.35), "CONCLUSIÓN", size=12, bold=True, color=ACCENT)
    txt(s, Inches(0.9), Inches(2.3), Inches(11.5), Inches(2.2),
        "La ventaja cuántica es real pero condicionada.\n"
        "Diagnosticar · remapear si procede ·\nactivar de forma selectiva e híbrida.",
        size=26, bold=True, color=WHITE)
    txt(s, Inches(0.9), Inches(5.2), Inches(11.5), Inches(0.7),
        "No sustituye al stack clásico. Lo complementa cuando la geometría o el QUBO lo justifican.",
        size=15, color=RGBColor(0xA8, 0xB0, 0xB8))
    rect(s, Inches(0.9), Inches(6.2), Inches(2.2), Inches(0.05), ACCENT)
    txt(s, Inches(0.9), Inches(6.45), Inches(11.5), Inches(0.4),
        "Gracias  ·  ¿Preguntas?", size=18, bold=True, color=WHITE)

    assert n == 20, n
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"Wrote {OUT} ({n} slides, {OUT.stat().st_size/1e6:.1f} MB)")


if __name__ == "__main__":
    build()
